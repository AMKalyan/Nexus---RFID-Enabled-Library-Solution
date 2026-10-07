"""Read rows from an uploaded .xlsx or .csv using the column mapping."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from app.core.errors import AppError

XLSX_MAGIC = b"PK\x03\x04"


@dataclass
class RawRow:
    row_no: int
    cells: dict[str, Any]  # keyed by Nexus field name


def detect_format(filename: str, content: bytes) -> str:
    name = filename.lower()
    if name.endswith(".xlsx") and content.startswith(XLSX_MAGIC):
        return "xlsx"
    if name.endswith(".csv") and b"\x00" not in content[:4096]:
        return "csv"
    raise AppError("UPLOAD_TYPE_INVALID", "Upload must be an .xlsx workbook or a .csv file", 422)


def _rows(fmt: str, content: bytes) -> Iterator[list[Any]]:
    if fmt == "xlsx":
        try:
            ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        except Exception as exc:  # noqa: BLE001 - any parse failure means an unusable file
            raise AppError("UPLOAD_UNREADABLE", "The workbook could not be read", 422) from exc
        yield from (list(r) for r in ws.iter_rows(values_only=True))
        return
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    yield from csv.reader(io.StringIO(text))


def read_rows(content: bytes, filename: str, mapping: dict[str, str]) -> Iterator[RawRow]:
    rows = _rows(detect_format(filename, content), content)
    try:
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
    except StopIteration as exc:
        raise AppError("UPLOAD_EMPTY", "The file has no header row", 422) from exc
    missing = [col for col in mapping if col not in header]
    if missing:
        raise AppError("UPLOAD_COLUMNS_MISSING", f"Missing columns: {', '.join(missing)}", 422)
    index = {mapping[col]: header.index(col) for col in mapping}
    row_no = 0
    for values in rows:
        if not any(v not in (None, "") for v in values):
            continue  # skip blank lines
        row_no += 1
        yield RawRow(row_no, {field: values[i] if i < len(values) else None for field, i in index.items()})
