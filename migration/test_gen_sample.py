import json

from gen_sample import HEADERS, generate
from openpyxl import load_workbook
from profile_sheet import profile


def test_fr10_generator_counts_and_manifest(tmp_path):
    path = tmp_path / "s.xlsx"
    m = generate(path, rows=1000, seed=1)
    rows = list(load_workbook(path, read_only=True).active.iter_rows(values_only=True))
    assert len(rows) == 1001 and list(rows[0]) == HEADERS
    assert 20 <= len(m.invalid_rows) <= 40
    assert 10 <= len(m.duplicate_rows) <= 30
    assert set(m.exact_duplicate_rows) | set(m.likely_duplicate_rows) == set(m.duplicate_rows)
    assert not set(m.invalid_rows) & set(m.duplicate_rows)
    saved = json.loads((tmp_path / "s.xlsx.manifest.json").read_text())
    assert saved["rows"] == 1000 and saved["invalid_rows"] == m.invalid_rows


def test_fr10_generator_is_deterministic(tmp_path):
    a = generate(tmp_path / "a.xlsx", rows=200, seed=7)
    b = generate(tmp_path / "b.xlsx", rows=200, seed=7)
    assert a.invalid_rows == b.invalid_rows and a.duplicate_rows == b.duplicate_rows


def test_fr10_every_defect_kind_is_injected(tmp_path):
    m = generate(tmp_path / "s.xlsx", rows=2000, seed=3)
    assert set(m.defects.values()) == {
        "TITLE_MISSING", "ISBN_CHECKSUM", "YEAR_RANGE", "ACCESSION_FORMAT", "ITEM_TYPE_UNKNOWN", "PRICE_NEGATIVE",
    }  # fmt: skip


def test_fr10_exact_duplicates_repeat_an_earlier_row(tmp_path):
    path = tmp_path / "s.xlsx"
    m = generate(path, rows=500, seed=5)
    rows = list(load_workbook(path, read_only=True).active.iter_rows(min_row=2, values_only=True))
    for row_no in m.exact_duplicate_rows:
        accession = rows[row_no - 1][0]
        assert any(r[0] == accession for r in rows[: row_no - 1])


def test_fr10_profile_reports_columns(tmp_path):
    path = tmp_path / "s.xlsx"
    generate(path, rows=300, seed=2)
    report = profile(path)
    assert report["rows"] == 300 and set(report["columns"]) == set(HEADERS)
    assert report["columns"]["Title"]["nulls"] >= 0 and report["columns"]["Accession No"]["distinct"] <= 300
